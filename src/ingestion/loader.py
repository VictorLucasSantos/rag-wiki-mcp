"""
src/ingestion/loader.py — Fase 1: Carregamento e parsing de documentos.

Responsabilidade:
  Receber arquivos em diferentes formatos e retornar texto
  com metadados estruturados (fonte, tipo, hash do conteúdo, etc.).

Por que isso importa:
  Um parser ruim (ex: extração ingênua de PDF com tabelas) indexa lixo.
  O modelo vai usar esse lixo para responder. As respostas vão parecer
  erradas "misteriosamente". Escolher o parser certo é decisão de qualidade.
"""

import hashlib
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterator


@dataclass
class Document:
    """
    Representa um documento carregado antes do chunking.

    Atributos:
        content  : texto extraído do arquivo
        metadata : informações sobre a origem do documento
    """
    content: str
    metadata: dict = field(default_factory=dict)

    def __repr__(self):
        preview = self.content[:80].replace("\n", " ")
        return f"Document(source={self.metadata.get('source', '?')}, preview='{preview}...')"


def content_hash(text: str) -> str:
    """Hash estável do conteúdo — usado na indexação incremental."""
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def source_id(path: Path) -> str:
    """
    Caminho canônico do arquivo ("docs/manual.pdf"): relativo ao diretório
    atual quando possível, para que "./docs" e "docs" gerem o mesmo ID.
    """
    resolved = path.resolve()
    try:
        return resolved.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return resolved.as_posix()


class _HTMLTextExtractor(HTMLParser):
    """Extrai texto visível de HTML, ignorando script/style/nav/footer."""

    SKIP = {"script", "style", "nav", "footer", "header", "noscript"}
    BLOCK = {"p", "div", "br", "li", "tr", "section", "article",
             "h1", "h2", "h3", "h4", "h5", "h6"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        elif tag in {"h1", "h2", "h3"} and not self._skip_depth:
            # Preserva a hierarquia como Markdown para o MarkdownChunker
            self.parts.append("\n\n" + "#" * int(tag[1]) + " ")
        elif tag in self.BLOCK:
            self.parts.append("\n\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)

    def text(self) -> str:
        return "".join(self.parts)


class DocumentLoader:
    """
    Carrega documentos de diferentes formatos.

    Suporte:
      .txt, .md     → leitura direta
      .html, .htm   → extração de texto visível (headers viram Markdown)
      .pdf          → pypdf (simples) — troque por unstructured/pymupdf
                      para PDFs com tabelas e layouts complexos
      .docx         → python-docx
      diretório     → carrega todos os arquivos suportados recursivamente

    Exemplo de uso:
        loader = DocumentLoader()
        docs = loader.load("./docs")           # carrega pasta inteira
        docs = loader.load("contrato.pdf")     # carrega arquivo único
    """

    SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".html", ".htm"}

    def load(self, path: str) -> list[Document]:
        """Ponto de entrada principal. Aceita arquivo ou diretório."""
        p = Path(path)

        if p.is_dir():
            return list(self._load_directory(p))
        elif p.is_file():
            return [self._load_file(p)]
        else:
            raise FileNotFoundError(f"Caminho não encontrado: {path}")

    def _load_directory(self, directory: Path) -> Iterator[Document]:
        """Carrega todos os arquivos suportados de um diretório."""
        for file_path in sorted(directory.rglob("*")):
            if file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                print(f"  📄 Carregando: {file_path.name}")
                yield self._load_file(file_path)

    def _load_file(self, file_path: Path) -> Document:
        """Direciona para o parser correto baseado na extensão."""
        ext = file_path.suffix.lower()

        if ext in {".txt", ".md"}:
            content, extra = file_path.read_text(encoding="utf-8"), {}
        elif ext in {".html", ".htm"}:
            content, extra = self._load_html(file_path), {}
        elif ext == ".pdf":
            content, extra = self._load_pdf(file_path)
        elif ext == ".docx":
            content, extra = self._load_docx(file_path), {}
        else:
            raise ValueError(f"Formato não suportado: {ext}")

        return Document(
            content=content,
            metadata={
                "source": source_id(file_path),
                "file_name": file_path.name,
                "file_type": ext.lstrip("."),
                "content_hash": content_hash(content),
                **extra,
            },
        )

    def _load_html(self, file_path: Path) -> str:
        parser = _HTMLTextExtractor()
        parser.feed(file_path.read_text(encoding="utf-8", errors="ignore"))
        return parser.text()

    def _load_pdf(self, file_path: Path) -> tuple[str, dict]:
        """
        Carrega PDFs.

        DECISÃO IMPORTANTE:
          Para PDFs simples (texto corrido) → pypdf é suficiente e mais rápido.
          Para PDFs com tabelas, imagens, layouts complexos → usar unstructured.io:

            from unstructured.partition.pdf import partition_pdf
            elements = partition_pdf(filename=str(file_path), strategy="hi_res")
            content = "\\n\\n".join([str(el) for el in elements])
        """
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("Instale pypdf: pip install pypdf")

        reader = PdfReader(str(file_path))
        pages = [
            f"[Página {i + 1}]\n{page.extract_text() or ''}"
            for i, page in enumerate(reader.pages)
        ]
        return "\n\n".join(pages), {"total_pages": len(reader.pages)}

    def _load_docx(self, file_path: Path) -> str:
        """Carrega arquivos .docx preservando parágrafos e títulos."""
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError("Instale python-docx: pip install python-docx")

        doc = DocxDocument(str(file_path))
        parts = []
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            # "Heading 1" / "Título 1" → "# ..." para preservar estrutura
            level = style.rsplit(" ", 1)[-1]
            if (style.startswith("heading") or style.startswith("título")) and level.isdigit():
                parts.append("#" * int(level) + " " + p.text)
            else:
                parts.append(p.text)
        return "\n\n".join(parts)
